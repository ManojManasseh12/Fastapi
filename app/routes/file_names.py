import io
from functools import lru_cache
from app.ai.rag.embed import models

from fastapi import APIRouter, UploadFile, File, HTTPException
import pymupdf
from pptx import Presentation
from starlette.concurrency import run_in_threadpool

router = APIRouter()

@lru_cache(maxsize=1)
def get_image_text():
    import easyocr
    return easyocr.Reader(["en"], gpu=False)

def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    filename_lower = filename.lower()
    text = ""

    if filename_lower.endswith(".pdf"):
        with pymupdf.open(stream=file_bytes, filetype="pdf") as pdf:
            for page in pdf:
                text += page.get_text()
    elif filename_lower.endswith(".pptx"):
        prs = Presentation(io.BytesIO(file_bytes))
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.text:
                            text += paragraph.text + "\n"
    elif filename_lower.endswith((".txt", ".md", ".csv", ".json")):
        text = file_bytes.decode("utf-8", errors="ignore")
    else:
        # Fallback: attempt PDF -> PPTX -> UTF-8 text
        try:
            with pymupdf.open(stream=file_bytes, filetype="pdf") as pdf:
                for page in pdf:
                    text += page.get_text()
        except Exception:
            try:
                prs = Presentation(io.BytesIO(file_bytes))
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for paragraph in shape.text_frame.paragraphs:
                                if paragraph.text:
                                    text += paragraph.text + "\n"
            except Exception:
                text = file_bytes.decode("utf-8", errors="ignore")

    return text

@router.post("/upload")
async def upload(file: UploadFile = File(...)):
    file_bytes = await file.read()
    filename = file.filename or ""

    try:
        text = extract_text_from_file(file_bytes, filename)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to process file: {str(e)}")

    words = text.split()
    if not words:
        return {
            "Total": 0,
            "Chunks": [],
            "Embeddings": []
        }

    chunks = []
    chunk_size, overlap = 200, 50
    step = chunk_size - overlap
    for start in range(0, len(words), step):
        chunk_words = words[start:start + chunk_size]
        chunks.append(" ".join(chunk_words))

        if start + chunk_size >= len(words):
            break  # reached the end of the text

    embeddings = models(chunks)
    if hasattr(embeddings, "tolist"):
        embeddings = embeddings.tolist()

    return {
        "Total": len(chunks),
        "Chunks": chunks,
        "Embeddings": embeddings
    }

@router.post("/upload-image")
async def upload_image(file: UploadFile = File(...)):
    image_bytes = await file.read()
    image_text = await run_in_threadpool(get_image_text)
    text = await run_in_threadpool(image_text.readtext, image_bytes, detail=0)
    return {"text": " ".join(text), "lines": text}
