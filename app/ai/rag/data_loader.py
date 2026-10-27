import io
from pptx import Presentation

def extract_pdf_text(file_bytes: bytes) -> str:
    """Extracts text from PDF bytes using PyMuPDF or pypdf fallback."""
    try:
        import pymupdf
        text = ""
        with pymupdf.open(stream=file_bytes, filetype="pdf") as pdf:
            for page in pdf:
                text += page.get_text()
        if text.strip():
            return text
    except Exception:
        pass

    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception:
        return ""

def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    """Extracts text based on the file extension (.pdf, .pptx, .txt, etc.)."""
    filename_lower = filename.lower()
    text = ""

    if filename_lower.endswith(".pdf"):
        text = extract_pdf_text(file_bytes)
    elif filename_lower.endswith(".pptx"):
        prs = Presentation(io.BytesIO(file_bytes))
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.text:
                            text += paragraph.text + "\n"
    elif filename_lower.endswith((".txt", ".md", ".csv", ".json")):
        text = file_bytes.decode("utf-8", errors="ignore")
    else:
        text = extract_pdf_text(file_bytes)
        if not text:
            try:
                prs = Presentation(io.BytesIO(file_bytes))
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for paragraph in shape.text_frame.paragraphs:
                                if paragraph.text:
                                    text += paragraph.text + "\n"
            except Exception:
                text = file_bytes.decode("utf-8", errors="ignore")

    return text

def chunk_text(text: str, chunk_size: int = 200, overlap: int = 50) -> list[str]:
    """Splits raw text into overlapping word chunks."""
    words = text.split()
    if not words:
        return []

    chunks = []
    step = chunk_size - overlap
    for start in range(0, len(words), step):
        chunk_words = words[start:start + chunk_size]
        chunks.append(" ".join(chunk_words))
        if start + chunk_size >= len(words):
            break

    return chunks
